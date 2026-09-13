"""The estate, as it actually is at this instant.

This is the module the whole reframe turns on. Before it, a triage tool printed
`disk 97%` because a random choice three milliseconds earlier had put 97 in a
dict. After it, the tool prints what the data box's disk is, because a collector
has been down for forty simulated minutes and the number has been climbing.

    One hidden cause drives both sides (TRIAGE.md §6). A scenario mutates this
    world; the telemetry the world emits and the readings the diagnostic tools
    take are both consequences of the same state. Nothing is scripted twice.

Held in memory, seeded once from the hub inventory sheet — the register is the
source of truth for the roster, telemetry is the source of truth for state, and
TELEMETRY-API.md §7 is emphatic that those are not the same claim.
"""

from __future__ import annotations

import random
from collections import deque
from dataclasses import dataclass, field
from datetime import UTC, datetime, timedelta
from pathlib import Path

# How long without a heartbeat before the cloud raises HUB-OFFLINE. A hub that
# is offline cannot report that it is offline (TELEMETRY-API.md §2), so this is
# the one code the world raises on the hub's behalf.
OFFLINE_THRESHOLD = timedelta(minutes=20)
HEARTBEAT_INTERVAL = timedelta(minutes=5)

# Per hub. Enough for fleet.heartbeat_gaps to show the run-up to a silence
# without holding twenty months of beats for 250 units.
HEARTBEAT_HISTORY = 240

ACCRUING = {"In service", "In service - monitored", "Loaner"}


@dataclass
class HubState:
    """One unit. Inventory columns, then everything the inventory cannot know."""

    hub_id: str
    serial: str | None
    customer_id: str | None
    organisation: str | None
    ward_id: str | None
    bed_id: str | None
    hw_revision: str | None
    sw_version: str | None
    install_date: datetime | None
    unit_status: str | None
    pairing_id: str | None
    pairing_status: str | None

    # Live state. Scenarios write here; tools and telemetry read here.
    last_seen: datetime | None = None
    battery_pct: float = 94.0
    on_dock: bool = True
    # The three layers a hub can fail at, held separately because the whole
    # point of OPS-SOP-004 is that from the cloud they look identical.
    socket_live: bool = True
    agent_up: bool = True
    link_up: bool = True
    socket_state: str = "live"
    pdu_port: str = "A07"
    restarts_24h: int = 0
    post_code: str | None = None
    rssi_dbm: int = -61
    wear_time_h: float = 12.0
    patch_lot: str | None = None
    bed_occupied: bool = True
    session_closed_cleanly: bool = False
    removal_reason: str | None = None
    clock_skew_s: int = 0
    reported_sw_version: str | None = (
        None  # diverges from inventory only under an off-taxonomy scenario
    )
    scenario_id: str | None = None
    offline_raised: bool = False

    heartbeats: deque = field(default_factory=lambda: deque(maxlen=HEARTBEAT_HISTORY))
    #: Last time each code was raised, so one fault is one event and not one
    #: per tick. A hub that raises CONN-LINK-LOSS every two minutes for a day
    #: has not had 720 incidents, and a register that says so is worse than no
    #: register at all.
    last_raised: dict = field(default_factory=dict)

    @property
    def accrues(self) -> bool:
        return self.unit_status in ACCRUING

    @property
    def powered(self) -> bool:
        """Mains, or what is left in the battery.

        A unit whose socket dies does not go dark that second — it runs down,
        raises BATT on the way, and then stops. That interval is the difference
        between a register row that says "lost power at 14:20" and one that can
        show the discharge curve underneath it.
        """
        return self.socket_live or self.battery_pct > 0

    @property
    def effective_sw(self) -> str | None:
        return self.reported_sw_version or self.sw_version

    def online(self, now: datetime) -> bool:
        return self.last_seen is not None and (now - self.last_seen) <= OFFLINE_THRESHOLD

    def inventory_row(self, now: datetime) -> dict:
        return {
            "hub_id": self.hub_id,
            "serial": self.serial,
            "customer_id": self.customer_id,
            "organisation": self.organisation,
            "ward_id": self.ward_id,
            "bed_id": self.bed_id,
            "hw_revision": self.hw_revision,
            "sw_version": self.effective_sw,
            "unit_status": self.unit_status,
            "pairing_id": self.pairing_id,
            "pairing_status": self.pairing_status,
            "install_date": self.install_date.date().isoformat() if self.install_date else None,
            "last_seen": self.last_seen.isoformat() if self.last_seen else None,
            "online": self.online(now),
            "battery_pct": round(self.battery_pct, 1),
            "accrues": self.accrues,
        }


@dataclass
class SiteState:
    """One customer's estate: the path to us, and the box in their comms room."""

    customer_id: str
    organisation: str | None

    # Network path
    tunnel_up: bool = True
    dns_ok: bool = True
    loss_pct: float = 0.0
    rtt_ms: int = 31
    change_window: str | None = None

    # Data box
    collector_active: bool = True
    queue_depth: int = 3
    disk_used_pct: int = 41
    cert_days_remaining: int = 121
    last_upload: datetime | None = None
    databox_agent: str = "2.4.1"

    # Written by registers.context when a scenario is an announced absence.
    planned_works_note: str | None = None

    @property
    def path_healthy(self) -> bool:
        return self.tunnel_up and self.dns_ok and self.loss_pct < 5

    @property
    def databox_healthy(self) -> bool:
        return self.collector_active and self.queue_depth < 500 and self.cert_days_remaining >= 0

    @property
    def host(self) -> str:
        return f"databox-{self.customer_id.lower().replace('cust-', '')}.pulseone.local"


class World:
    """The fleet. Seeded from the inventory, moved by scenarios, read by everybody."""

    def __init__(self, *, seed: int = 20260913) -> None:
        self.hubs: dict[str, HubState] = {}
        self.sites: dict[str, SiteState] = {}
        self.rng = random.Random(seed)
        self.events: deque = deque(maxlen=20000)
        self._seq = 0

    # -- seeding -----------------------------------------------------------

    @classmethod
    def from_inventory(cls, workbook: Path, *, now: datetime, seed: int = 20260913) -> World:
        """250 rows off the `PulseOne Hub Inventory` sheet, then made live.

        `Last Online (local)` in the sheet is a stale column by construction —
        it is what the spreadsheet last happened to record. The world does not
        trust it: every in-service unit starts heartbeating now, because the
        simulation begins with a fleet that is up.
        """
        import openpyxl

        world = cls(seed=seed)
        wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)

        orgs = {}
        ws = wb["Organisations"]
        rows = ws.iter_rows(values_only=True)
        header = [str(c) if c else "" for c in next(rows)]
        for row in rows:
            r = dict(zip(header, row, strict=False))
            if r.get("CustomerID"):
                orgs[r["CustomerID"]] = r.get("Organisation Name")

        ws = wb["PulseOne Hub Inventory"]
        rows = ws.iter_rows(values_only=True)
        header = [str(c) if c else "" for c in next(rows)]
        for row in rows:
            r = dict(zip(header, row, strict=False))
            hub_id = r.get("HubID")
            if not hub_id:
                continue
            customer_id = r.get("CustomerID")
            hub = HubState(
                hub_id=hub_id,
                serial=r.get("Serial Number"),
                customer_id=customer_id,
                organisation=r.get("Organisation") or orgs.get(customer_id),
                ward_id=r.get("WardID"),
                bed_id=r.get("BedID"),
                hw_revision=r.get("HW Revision"),
                sw_version=r.get("SW Version (last observed)"),
                install_date=_as_dt(r.get("Install Date")),
                unit_status=r.get("Unit Status"),
                pairing_id=r.get("Pairing ID"),
                pairing_status=r.get("Pairing Status"),
            )
            hub.pdu_port = f"A{world.rng.randint(1, 24):02d}"
            hub.battery_pct = round(world.rng.uniform(72, 99), 1)
            hub.rssi_dbm = -world.rng.randint(52, 72)
            hub.wear_time_h = round(world.rng.uniform(2, 60), 1)
            # A unit that does not accrue is not on a ward, so there is nobody
            # in the bed and nothing to triage (estate.py has always said so).
            hub.bed_occupied = hub.accrues
            if hub.accrues:
                hub.last_seen = now
            world.hubs[hub_id] = hub
            if customer_id and customer_id not in world.sites:
                world.sites[customer_id] = SiteState(
                    customer_id=customer_id, organisation=orgs.get(customer_id)
                )
                world.sites[customer_id].last_upload = now

        wb.close()
        return world

    # -- lookups -----------------------------------------------------------

    def site_of(self, hub: HubState) -> SiteState | None:
        return self.sites.get(hub.customer_id) if hub.customer_id else None

    def reachable(self, hub: HubState) -> bool:
        """Whether an ssh session to the unit would establish.

        A hub with no power and a hub behind a dead tunnel both fail to answer,
        and from the bastion they are indistinguishable — which is exactly why
        OPS-SOP-004 checks the socket and the path before it blames the unit.
        The data box is deliberately *not* in here: a box that has stopped
        shipping still passes traffic, so the unit answers and looks perfect.
        """
        site = self.site_of(hub)
        return hub.powered and (site is None or site.path_healthy)

    def visible_to_cloud(self, hub: HubState) -> bool:
        """Whether a heartbeat from this unit would reach us.

        Derived through the stack rather than set by a scenario, so a silence
        always has a layer under it. This is the single function that makes
        `data-box-backlog` and `power-loss` present identically from the cloud,
        which is the confusion the procedure exists to resolve.
        """
        if not (hub.powered and hub.agent_up):
            return False
        site = self.site_of(hub)
        return site is None or (site.path_healthy and site.databox_healthy)

    def hubs_at(self, customer_id: str) -> list[HubState]:
        return [h for h in self.hubs.values() if h.customer_id == customer_id]

    def live_hubs(self) -> list[HubState]:
        return [h for h in self.hubs.values() if h.accrues]

    # -- the stream --------------------------------------------------------

    def emit(
        self,
        hub: HubState,
        code: str,
        now: datetime,
        *,
        severity: str = "warn",
        payload: dict | None = None,
        scenario_id: str | None = None,
    ) -> dict:
        """Append one telemetry event, in the shape events.ndjson already has.

        `ts_device` carries the hub's own clock skew and `ts_received` is UTC at
        the ingest edge. They disagree, and which one you window on changes the
        answer — a property the frozen file has and the live stream must keep,
        because an off-taxonomy scenario makes the gap enormous on purpose.
        """
        self._seq += 1
        ts_device = now + timedelta(seconds=hub.clock_skew_s)
        event = {
            "event_id": f"EVT-SIM-{self._seq:06d}",
            "hub_id": hub.hub_id,
            "serial": hub.serial,
            "customer_id": hub.customer_id,
            "ward_id": hub.ward_id,
            "bed_id": hub.bed_id,
            "ts_device": ts_device.isoformat(),
            "ts_received": now.isoformat(),
            "code": code,
            "severity": severity,
            "sw_version": hub.effective_sw,
            "hw_revision": hub.hw_revision,
            "source": "device" if code != "HUB-OFFLINE" else "cloud",
            "dedupe_key": f"{hub.hub_id}:{code}:{ts_device:%Y-%m-%dT%H:%M}",
            "payload": payload or {},
            "scenario_id": scenario_id,
        }
        self.events.append(event)
        return event

    def heartbeat(self, hub: HubState, now: datetime) -> None:
        hub.last_seen = now
        hub.offline_raised = False
        hub.heartbeats.append(
            {
                "ts": now.isoformat(),
                "battery_pct": round(hub.battery_pct, 1),
                "sw_version": hub.effective_sw,
                "pairing_status": hub.pairing_status,
                "rssi_dbm": hub.rssi_dbm,
            }
        )

    def since(self, cursor: int) -> tuple[list[dict], int]:
        """Events after `cursor`, and the new cursor. The poller's whole API."""
        out = [e for e in self.events if _seq_of(e) > cursor]
        return out, (_seq_of(out[-1]) if out else cursor)


def _seq_of(event: dict) -> int:
    try:
        return int(event["event_id"].rsplit("-", 1)[-1])
    except (KeyError, ValueError):
        return 0


def _as_dt(value) -> datetime | None:
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=UTC)
    return None
